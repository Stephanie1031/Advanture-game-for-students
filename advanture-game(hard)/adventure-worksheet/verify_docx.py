#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""校园冒险作业纸自检 / adventure worksheet self-check

    python3 adventure-worksheet/verify_docx.py

检查 / checks:
  1. docx 结构完好（zip + 所有 XML 合法）
  2. 每一段中文都指定了中文字体（否则某些电脑上会变方框）
  3. 分页：每一页内容高度都在 A4 可打印区域内
  4. **作业纸印的预期输出，确实来自 classes.py 的 docstring**（逐条比对）
  5. 上级目录的源文件没有被改动（用 git / 哈希无法判断，这里改为检查文件清单）

第 4 项最重要：这张纸本来就替代 ok 测试，纸上印的输出必须和题面一致。
"""

import io
import os
import re
import sys
import zipfile
import xml.dom.minidom

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:                                     # pragma: no cover
    sys.exit('请先安装 python-docx / please install python-docx')

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(HERE)
WORKSHEET = os.path.join(HERE, '校园冒险-作业纸.docx')
ANSWERS = os.path.join(HERE, '校园冒险-参考答案.docx')
EMU_PER_CM = 360000.0

sys.path.insert(0, '/tmp')                              # reuse the layout estimator
failures = []


def check(condition, label):
    print(('  PASS  ' if condition else '  FAIL  ') + label)
    if not condition:
        failures.append(label)


def doc_text(path):
    doc = Document(path)
    lines = []
    for p in doc.element.body.findall('.//' + qn('w:p')):
        lines.append(''.join(t.text or '' for t in p.findall('.//' + qn('w:t'))))
    return '\n'.join(lines)


def check_structure(path, label):
    with zipfile.ZipFile(path) as archive:
        bad = []
        for item in archive.namelist():
            if item.endswith('.xml') or item.endswith('.rels'):
                try:
                    xml.dom.minidom.parseString(archive.read(item))
                except Exception as error:
                    bad.append((item, error))
        check(not bad, '{} 结构完好（{} 个部件 XML 合法）'.format(label, len(archive.namelist())))


def check_fonts(path, label):
    doc = Document(path)
    missing = 0
    total = 0
    for run in doc.element.body.findall('.//' + qn('w:r')):
        text = ''.join(t.text or '' for t in run.findall(qn('w:t')))
        if not text.strip():
            continue
        total += 1
        rpr = run.find(qn('w:rPr'))
        fonts = rpr.find(qn('w:rFonts')) if rpr is not None else None
        if fonts is None or fonts.get(qn('w:eastAsia')) is None:
            missing += 1
    check(missing == 0, '{} 的 {} 段文字都指定了中文字体'.format(label, total))


def check_pages(path, label):
    """Recompute page heights with the shared honest model."""
    import estimate_layout as est
    doc = Document(path)
    sec = doc.sections[0]
    usable_h = (sec.page_height - sec.top_margin - sec.bottom_margin) / EMU_PER_CM
    usable_w = (sec.page_width - sec.left_margin - sec.right_margin) / EMU_PER_CM
    heights, used = [], 0.0
    for child in doc.element.body.iterchildren():
        tag = child.tag.split('}')[1]
        if tag == 'p':
            if any(br.get(qn('w:type')) == 'page'
                   for br in child.findall('.//' + qn('w:br'))):
                heights.append(used)
                used = 0.0
                continue
            used += est.para_height(child, usable_w)
        elif tag == 'tbl':
            used += est.table_height(child, usable_w)
    heights.append(used)
    empty = [i + 1 for i, h in enumerate(heights) if h < 1.0]
    worst = max(heights)
    check(worst <= usable_h,
          '{} {} 页都在 A4 可打印范围内（最高 {:.1f} / {:.1f} cm）'.format(
              label, len(heights), worst, usable_h))
    check(not empty, '{} 没有空白页'.format(label))
    print('         每页估计高度 / heights: '
          + ', '.join('p{}={:.1f}'.format(i + 1, h) for i, h in enumerate(heights)))


def check_outputs_match_docstrings():
    """作业纸里印的预期输出，必须能在 classes.py 的 docstring 里找到。"""
    with io.open(os.path.join(PROJECT, 'classes.py'), encoding='utf-8') as handle:
        source = handle.read()

    # 每个方法的 docstring 范围
    def docstring_of(method):
        match = re.search(r'def {}\(.*?\):\s*"""(.*?)"""'.format(method),
                          source, re.S)
        return match.group(1) if match else ''

    docstrings = {name: docstring_of(name)
                  for name in ('go_to', 'talk_to', 'take', 'unlock')}
    for name, text in docstrings.items():
        check(bool(text), 'classes.py 里能找到 {} 的 docstring'.format(name))

    worksheet = doc_text(WORKSHEET)

    # 作业纸里那些「来自 docstring」的输出行
    expected_lines = [
        'You are at GBC, Golden Bear Cafe.',
        "Can't go to GBC from GBC.",
        'Try looking around to see where to go.',
        'Sather Gate is locked! Go look for a key to unlock it',
        'Person has to be a string.',
        'Jerry says: I am not the Jerry you are looking for.',
        'Tiffany is not here.',
        'Thing should be a string.',
        'orange is not here.',
        'Player takes the Lemon',
        'Lemon is not here.',
        "FSM can't be unlocked without a key!",
        'FSM is now unlocked!',
        'FSM is already unlocked!',
        'Player takes the SkeletonKey',
    ]
    for line in expected_lines:
        check(line in worksheet, '作业纸印了：{}'.format(line[:46]))

    merged = '\n'.join(docstrings.values())
    missing = [line for line in expected_lines if line not in merged]
    check(not missing,
          '以上输出全部来自 docstring（未在 docstring 中出现的：{}）'.format(missing or '无'))

    # 正确答案必须和题面一致
    solution_lines = [
        'You are at {}, {}',
        'self.place = destination_place',
        "print('{} says: {}'.format(person, character.talk()))",
        'self.place.take(thing)',
        'place.locked = False',
        'key.use(next_place)',
    ]
    answers = doc_text(ANSWERS)
    for line in solution_lines:
        check(line in answers, '参考答案含：{}'.format(line[:46]))


def check_given_code_matches_source():
    """「已经给你的」那些代码，必须和 classes.py 里的原文一致。"""
    with io.open(os.path.join(PROJECT, 'classes.py'), encoding='utf-8') as handle:
        source = handle.read()
    worksheet = doc_text(WORKSHEET)

    given = [
        'destination_place = self.place.get_neighbor(location)',
        "print(destination_place.name, 'is locked! Go look for a key to unlock it')",
        "print('Person has to be a string.')",
        "print('Thing should be a string.')",
        "print(\"Place must be a string\")",
        'if type(item) == Key:',
    ]
    for line in given:
        check(line in source and line in worksheet,
              '「已给你的代码」与 classes.py 一致：{}'.format(line[:44]))

    check('""" Implement Key here! """' in worksheet,
          '作业纸标出了 Key 要写在哪（""" Implement Key here! """）')
    check('class Thing(object):' in worksheet and 'class Thing(object):' in source,
          '作业纸上的 Thing 父类与 classes.py 一致')


def check_sources_untouched():
    """确认工作只发生在本文件夹里 / nothing outside this folder was created."""
    known = {'classes.py', 'data.py', 'adventure.py', 'adventure_gui.py',
             'README.md', 'Questions.md', 'ok', 'adventure.exe',
             '.ok_history', '.ok_storage.db', '.DS_Store', '__pycache__',
             'adventure-worksheet', 'dino-pet-game'}
    extra = sorted(name for name in os.listdir(PROJECT) if name not in known)
    check(not extra, '上级目录没有新增文件（多余：{}）'.format(extra or '无'))
    with io.open(os.path.join(PROJECT, 'classes.py'), encoding='utf-8') as handle:
        original = handle.read()
    check('*** YOUR CODE HERE ***' in original,
          'classes.py 仍是待完成的原状（含 YOUR CODE HERE 占位）')
    check('""" Implement Key here! """' in original,
          'classes.py 里的 Key 占位注释没有被改动')


def main():
    print('=' * 70)
    print('校园冒险作业纸自检 / adventure worksheet self-check')
    print('=' * 70)

    check_structure(WORKSHEET, '作业纸')
    check_fonts(WORKSHEET, '作业纸')
    check_pages(WORKSHEET, '作业纸')
    print()
    check_structure(ANSWERS, '参考答案')
    check_fonts(ANSWERS, '参考答案')
    check_pages(ANSWERS, '参考答案')
    print()
    check_outputs_match_docstrings()
    print()
    check_given_code_matches_source()
    print()
    check_sources_untouched()

    print('\n' + '=' * 70)
    if failures:
        print('{} 项未通过 / {} failed:'.format(len(failures), len(failures)))
        for item in failures:
            print('   - ' + item)
        return 1
    print('全部通过 / all checks passed')
    return 0


if __name__ == '__main__':
    sys.exit(main())
